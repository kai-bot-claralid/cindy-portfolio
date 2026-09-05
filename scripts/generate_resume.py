from pathlib import Path
import hashlib

# Compatibility for the system OpenSSL/Python combination used during generation.
_original_md5 = hashlib.md5
def _compatible_md5(*args, **kwargs):
    kwargs.pop("usedforsecurity", None)
    return _original_md5(*args, **kwargs)
hashlib.md5 = _compatible_md5

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "resume"
OUT.mkdir(exist_ok=True)

NAME = "Cindy Guzmán"
TITLE = "Web Developer & UX/UI Designer"
CONTACT = "Santo Domingo, Dominican Republic | cindyclaralid@gmail.com | claralid.com | linkedin.com/in/claralid"

SUMMARY = (
    "Web Developer and UX/UI Designer with experience designing, building and maintaining responsive websites, "
    "digital platforms and product interfaces. Strong background in WordPress and CMS-based development, supported "
    "by UX research, information architecture, prototyping and front-end fundamentals. Uses AI-assisted development "
    "workflows to move product requirements and interface designs toward tested, functional releases."
)

SKILLS = [
    ("Web Development & CMS", "WordPress, Joomla, Shopify, PrestaShop, DSpace, HTML, CSS, basic JavaScript, responsive implementation, website maintenance, deployment"),
    ("UX/UI & Product Design", "UX research, user flows, information architecture, wireframing, interface design, prototyping, design systems, usability, accessibility considerations"),
    ("Development Workflow", "Git (basic), GitHub, Visual Studio Code, Cursor, Claude Code, Codex, Antigravity, AI-assisted debugging and testing"),
    ("Design Tools", "Figma, Framer, Adobe Photoshop, Adobe Illustrator, Affinity, Canva"),
]

EXPERIENCE = [
    ("Content & Publications", "UNPHU", "Current", "Support content, publications and digital work across institutional platforms in a university environment."),
    ("Web Developer", "Behealth PR", "Jun 2022 – Present", "Build, update and maintain production websites for a Puerto Rico-based healthcare client, improving content delivery, usability and day-to-day reliability."),
    ("UX/UI Designer", "Cervecería Nacional Dominicana", "Sep 2022 – Feb 2023", "Translated business and user needs into clear, usable interfaces while collaborating with a multidisciplinary, multinational team."),
    ("Web Developer & Web Department Coordinator", "Gmedia Dominicana", "Aug 2019 – Nov 2021", "Coordinated web production and built and maintained WordPress, Joomla and Shopify websites for multiple clients."),
]

PROJECTS = [
    ("Amino | Digital Product / Web Platform", "Product Strategy, UX Research, UX/UI, AI-Assisted Development, Implementation", "Designed, built and launched a functional platform for creating microsites and link-based digital spaces; managed the process from product definition and interface design through testing, deployment and iteration. | aminoweb.com"),
    ("MinoSteps | Accessible Digital Product", "Product Research, UX/UI, Accessibility, Web Development", "Designed and developed a digital experience that makes everyday guidance calmer, clearer and easier to follow, using accessibility thinking and user-centered product decisions. | minosteps.com"),
    ("UNPHU | Institutional Website", "UX/UI Research, Interface Design, Collaboration", "Contributed research-informed interface work within an established institutional ecosystem and multidisciplinary team. | unphu.edu.do"),
]

EDUCATION = "Multimedia Technology | Instituto Tecnológico de las Américas (ITLA)"


def build_docx():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(9.5)
    styles["Normal"].paragraph_format.space_after = Pt(3)
    styles["Normal"].paragraph_format.line_spacing = 1.05

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(NAME)
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(20)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(TITLE)
    r.bold = True
    r.font.size = Pt(11)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(CONTACT)

    def heading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text.upper())
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(70, 77, 58)

    heading("Professional Summary")
    doc.add_paragraph(SUMMARY)

    heading("Core Skills")
    for label, value in SKILLS:
        p = doc.add_paragraph(style=None)
        p.paragraph_format.left_indent = Inches(0.12)
        p.paragraph_format.first_line_indent = Inches(-0.12)
        p.add_run(f"{label}: ").bold = True
        p.add_run(value)

    heading("Professional Experience")
    for role, company, dates, detail in EXPERIENCE:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.add_run(f"{role} | {company}").bold = True
        p.add_run(f" | {dates}")
        p = doc.add_paragraph(detail)
        p.paragraph_format.left_indent = Inches(0.12)

    heading("Selected Projects")
    for name, role, detail in PROJECTS:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.add_run(name).bold = True
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.12)
        p.add_run("Role: ").bold = True
        p.add_run(role)
        p = doc.add_paragraph(detail)
        p.paragraph_format.left_indent = Inches(0.12)

    heading("Education")
    doc.add_paragraph(EDUCATION)
    doc.save(OUT / "Cindy-Guzman-ATS-Resume.docx")


def build_pdf():
    out = OUT / "Cindy-Guzman-ATS-Resume.pdf"
    doc = SimpleDocTemplate(
        str(out), pagesize=LETTER,
        rightMargin=0.58 * inch, leftMargin=0.58 * inch,
        topMargin=0.48 * inch, bottomMargin=0.48 * inch,
        title=f"{NAME} — ATS Resume", author=NAME,
    )
    olive = colors.HexColor("#4B503E")
    styles = getSampleStyleSheet()
    name_style = ParagraphStyle("Name", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=20, leading=22, alignment=TA_CENTER, spaceAfter=2)
    title_style = ParagraphStyle("Title", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=10.5, leading=13, alignment=TA_CENTER, spaceAfter=2)
    contact_style = ParagraphStyle("Contact", parent=styles["Normal"], fontName="Helvetica", fontSize=8.2, leading=11, alignment=TA_CENTER, spaceAfter=7)
    section_style = ParagraphStyle("Section", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=9.5, leading=12, textColor=olive, spaceBefore=6, spaceAfter=2, borderWidth=0, uppercase=True)
    body = ParagraphStyle("Body", parent=styles["BodyText"], fontName="Helvetica", fontSize=8.5, leading=10.8, spaceAfter=2)
    indented = ParagraphStyle("Indented", parent=body, leftIndent=9)

    story = [Paragraph(NAME, name_style), Paragraph(TITLE, title_style), Paragraph(CONTACT, contact_style)]

    def heading(text):
        story.append(Paragraph(text.upper(), section_style))

    heading("Professional Summary")
    story.append(Paragraph(SUMMARY, body))
    heading("Core Skills")
    for label, value in SKILLS:
        story.append(Paragraph(f"<b>{label}:</b> {value}", indented))
    heading("Professional Experience")
    for role, company, dates, detail in EXPERIENCE:
        story.append(Paragraph(f"<b>{role} | {company}</b> | {dates}", body))
        story.append(Paragraph(detail, indented))
    heading("Selected Projects")
    for name, role, detail in PROJECTS:
        story.append(Paragraph(f"<b>{name}</b>", body))
        story.append(Paragraph(f"<b>Role:</b> {role}", indented))
        story.append(Paragraph(detail, indented))
    heading("Education")
    story.append(Paragraph(EDUCATION, body))
    doc.build(story)


if __name__ == "__main__":
    build_docx()
    build_pdf()
